from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn


DOCUMENT_DIR = Path(r"C:\mic_practice\output\documents")
SOURCE = DOCUMENT_DIR / (
    "Марикин_А_Н_и_др_Исследование_прямого_нейроуправления_27_08_2026_"
    "скорость_момент_уточнено.docx"
)
OUTPUT = DOCUMENT_DIR / (
    "Марикин_А_Н_и_др_Исследование_прямого_нейроуправления_27_08_2026_"
    "скорость_момент_сокращено.docx"
)


REPLACEMENTS = {
    6: (
        "Цель: Оценить возможность прямого нейросетевого управления "
        "асинхронным электроприводом без полной предварительной идентификации "
        "конкретного двигателя и сопоставить MIC с FOC при постоянном задании "
        "скорости и изменении момента нагрузки."
    ),
    8: (
        "Практическая значимость: Политику предполагается адаптировать к "
        "конкретному двигателю по короткому окну измерений и переносить на "
        "встроенную платформу, сокращая объём предварительной параметрической "
        "идентификации."
    ),
    12: (
        "Научная новизна состоит в прямом формировании RL-агентом воздействий "
        "в d-q координатах без токового каскада PI-регуляторов. Политика "
        "обучается на семействе цифровых моделей и адаптируется по короткому "
        "окну измерений, поэтому рабочий контур не требует полной идентификации "
        "конкретного двигателя."
    ),
    13: (
        "В отличие от DeepMPC, использующего обучаемую модель в предиктивном "
        "контуре [6], MIC непосредственно формирует воздействия на инвертор. "
        "Агент обучается методом проб и ошибок по функции вознаграждения [4], "
        "что сокращает ручную настройку."
    ),
    25: (
        "В MIC RL-агент по измеряемым скорости и токам формирует воздействия "
        "на инвертор. Цифровой двойник используется при обучении, а короткое "
        "окно измерений — при адаптации к конкретному приводу. После сжатия "
        "политика исполняется на встроенной платформе внутри независимых "
        "ограничений по току, напряжению, температуре и состояниям силовых ключей."
    ),
    28: (
        "Сравнение MIC и FOC выполнено при неизменном задании скорости и двух "
        "значениях момента нагрузки: 0,5Mном (стадия 0) и 1,0Mном (стадия 1). "
        "Оценивались ток статора, активная мощность и ошибка скорости."
    ),
    36: (
        "На рисунке 3 стадии 0 и 1 соответствуют нагрузкам 0,5Mном и 1,0Mном "
        "при одинаковом задании скорости. В первом режиме стратегии практически "
        "совпадают. Во втором MIC повышает скорость и выходную мощность на "
        "10,4 %, но также ток на 7,6 %, входную мощность на 13,3 % и снижает "
        "КПД на 2,0 процентного пункта. Поэтому безусловный энергетический "
        "выигрыш не подтверждён."
    ),
    38: (
        "1. MIC формирует воздействия в d-q координатах; цифровой двойник "
        "используется только при обучении."
    ),
    39: (
        "2. При 0,5Mном различий почти нет; при 1,0Mном ток MIC выше на "
        "7,6 %, поэтому его снижение не подтверждено."
    ),
    40: (
        "3. Снижение мощности на 8,4 % относится только к интервалу рисунка 2 "
        "и не доказывает общую экономию."
    ),
    41: (
        "4. Энергетический эффект требует одинаковых скорости, момента и "
        "выходной мощности, а также повторных опытов."
    ),
    42: (
        "5. Подход работоспособен и готов к стендовой адаптации с независимой "
        "защитой."
    ),
    67: (
        "Summary\n"
        "Objective: To compare MIC with FOC at a constant speed reference under "
        "changing load torque without complete prior identification of the "
        "particular motor. Methods: A digital twin was used; current, active "
        "power, and speed error were evaluated at 0.5 and 1.0 rated torque. "
        "Practical importance: The policy can be adapted from a short measurement "
        "window and deployed on an embedded platform."
    ),
}


EXPECTED_PREFIXES = {
    6: "Цель:",
    8: "Практическая значимость:",
    12: "Научная новизна",
    13: "Применение алгоритмов",
    25: "В системе MIC",
    28: "Для оценки поведения",
    36: "На рисунке 3",
    38: "1.",
    39: "2.",
    40: "3.",
    41: "4.",
    42: "5.",
    67: "Summary",
}


SEGMENTED_REPLACEMENTS = {
    6: [
        ("Цель: Оценить возможность прямого нейросетевого управления асинхронным электроприводом без полной предварительной идентификации конкретного двигателя и сопоставить MIC с FOC ", False),
        ("при постоянном задании скорости и изменении момента нагрузки.", True),
    ],
    8: [
        ("Практическая значимость: Политику предполагается адаптировать к конкретному двигателю ", False),
        ("по короткому окну измерений", True),
        (" и переносить на встроенную платформу, сокращая объём предварительной параметрической идентификации.", False),
    ],
    12: [
        ("Научная новизна состоит в прямом формировании RL-агентом воздействий в d-q координатах без токового каскада PI-регуляторов. ", False),
        ("Политика обучается на семействе цифровых моделей и адаптируется по короткому окну измерений", True),
        (", поэтому рабочий контур не требует полной идентификации конкретного двигателя.", False),
    ],
    25: [
        ("В MIC RL-агент по измеряемым скорости и токам формирует воздействия на инвертор. ", False),
        ("Цифровой двойник используется при обучении, а короткое окно измерений — при адаптации к конкретному приводу.", True),
        (" После сжатия политика исполняется на встроенной платформе внутри независимых ограничений по току, напряжению, температуре и состояниям силовых ключей.", False),
    ],
    67: [
        ("Summary\nObjective: To compare MIC with FOC ", False),
        ("at a constant speed reference under changing load torque", True),
        (" without complete prior identification of the particular motor. Methods: A digital twin was used; current, active power, and speed error were evaluated at 0.5 and 1.0 rated torque. Practical importance: The policy can be adapted from ", False),
        ("a short measurement window", True),
        (" and deployed on an embedded platform.", False),
    ],
}


def replace_paragraph_text(paragraph, text):
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_rpr = deepcopy(paragraph.runs[0]._r.rPr)

    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)

    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.highlight_color = WD_COLOR_INDEX.RED


def replace_paragraph_segments(paragraph, segments):
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_rpr = deepcopy(paragraph.runs[0]._r.rPr)

    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    for text, highlighted in segments:
        run = paragraph.add_run(text)
        if first_rpr is not None:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, deepcopy(first_rpr))
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.highlight_color = WD_COLOR_INDEX.RED if highlighted else None


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def highlighted_word_count(document):
    count = 0
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.highlight_color == WD_COLOR_INDEX.RED:
                count += len(run.text.split())
    return count


def main():
    document = Document(SOURCE)
    paragraphs = document.paragraphs
    before = highlighted_word_count(document)

    for index, replacement in REPLACEMENTS.items():
        paragraph = paragraphs[index]
        expected = EXPECTED_PREFIXES[index]
        if not paragraph.text.startswith(expected):
            raise RuntimeError(
                f"Paragraph {index} does not start with {expected!r}: {paragraph.text!r}"
            )
        if index in SEGMENTED_REPLACEMENTS:
            replace_paragraph_segments(paragraph, SEGMENTED_REPLACEMENTS[index])
        else:
            replace_paragraph_text(paragraph, replacement)

    removals = {
        26: "Принципиально важно",
        27: "После завершения этапа обучения",
        43: "Таким образом, моделирование",
    }
    for index, expected in removals.items():
        paragraph = paragraphs[index]
        if not paragraph.text.startswith(expected):
            raise RuntimeError(f"Unexpected paragraph {index}: {paragraph.text!r}")
        remove_paragraph(paragraph)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    check = Document(OUTPUT)
    after = highlighted_word_count(check)
    print(f"Saved: {OUTPUT}")
    print(f"Red-highlighted words: {before} -> {after}")
    print(f"Paragraphs: {len(paragraphs)} -> {len(check.paragraphs)}")


if __name__ == "__main__":
    main()
